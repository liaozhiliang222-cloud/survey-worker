import json,sys,zipfile
from pathlib import Path
from io import BytesIO
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from pptx import Presentation
from openpyxl import load_workbook
from pptx_report.research_delivery import build_package
out=Path('.data/s3-delivery')
for name in ['quantitative','qualitative','mixed']:
    payload=json.loads((out/f'{name}.json').read_text(encoding='utf-8-sig'))
    result=build_package(payload);(out/f'{name}.zip').write_bytes(result['content']);(out/f'{name}.pptx').write_bytes(result['pptx'])
    prs=Presentation(BytesIO(result['pptx']));assert len(prs.slides)==result['slide_count']
    charts=[shape.chart for slide in prs.slides for shape in slide.shapes if shape.has_chart]
    if name!='qualitative':assert tuple(charts[0].series[0].values)==(0.0,100.0)
    texts='\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    if name!='quantitative':assert '包装上的说明让我更容易理解这款产品。' in texts
    z=zipfile.ZipFile(BytesIO(result['content']));assert len(z.namelist())==5
    wb=load_workbook(BytesIO(z.read('分析结果与覆盖.xlsx')))
    if name!='qualitative':assert wb['交付数据']['D2'].value==0 and wb['交付数据']['D3'].value==100
    for slide in prs.slides:
        for s in slide.shapes:assert s.left>=0 and s.top>=0 and s.left+s.width<=prs.slide_width and s.top+s.height<=prs.slide_height
    print(name,result['slide_count'],'slides: native objects and workbook readback passed')
# Long title, exact quote continuation and dense body remain intact.
payload=json.loads((out/'mixed.json').read_text(encoding='utf-8-sig'));page=payload['script']['pages'][1]
page['title']='信息解释的完整性影响理解与选择，报告需要呈现长标题、长原声及密集信息，并保留每项证据的真实来源与解释边界'
page['quotes'][0]['text']='受访者明确说明包装信息很重要。'*40
payload['records'][2]['evidence']['value']['quote']=page['quotes'][0]['text']
page['content_structure']=[{'title':'解释边界','body':'仅用于描述当前访谈，不推断总体因果。'*35}]
r=build_package(payload);(out/'dense.pptx').write_bytes(r['pptx'])
p=Presentation(BytesIO(r['pptx']));assert len(p.slides)>4
for kind in ['bad_number','bad_quote','unsupported','missing']:
    value=json.loads((out/'mixed.json').read_text(encoding='utf-8-sig'))
    if kind=='bad_number':value['script']['pages'][0]['data_points'][0]['value']='1'
    if kind=='bad_quote':value['script']['pages'][1]['quotes'][0]['text']='捏造的原声'
    if kind=='unsupported':value['script']['pages'][0]['page_type']='unknown'
    if kind=='missing':value['script']['pages'][0]['data_points'][0]['value']=None
    try:build_package(value)
    except ValueError:pass
    else:raise AssertionError(kind)
print('Long content continuation and invalid data rejection passed')

# Exercise the real HTTP boundary (project scope + ZIP response + invalid data).
from fastapi.testclient import TestClient
from deploy.aliyun_api import app
client=TestClient(app)
payload=json.loads((out/'mixed.json').read_text(encoding='utf-8'))
response=client.post('/api/pptx-report/research-report-package',headers={'X-Project-Id':'p'},json=payload)
assert response.status_code==200,response.text[:500]
assert response.content[:2]==b'PK'
assert client.post('/api/pptx-report/research-report-package',headers={'X-Project-Id':'other'},json=payload).status_code==403
payload['script']['pages'][0]['data_points'][0]['unit']='NPS'
assert client.post('/api/pptx-report/research-report-package',headers={'X-Project-Id':'p'},json=payload).status_code==400
print('FastAPI scope, ZIP download and unit mismatch rejection passed')

# Weighted/mean units and raw precision are preserved, without inferred scaling.
payload=json.loads((out/'quantitative.json').read_text(encoding='utf-8'))
point=payload['script']['pages'][0]['data_points'][0]
point.update(value='0.5',unit='均值')
payload['records'][0]['evidence']['value'].update(value=0.5,unit='均值',weighted=True,weighted_base=19.5)
payload['records'][0]['dataset']['type']='weighted'
r=build_package(payload);z=zipfile.ZipFile(BytesIO(r['content']));wb=load_workbook(BytesIO(z.read('分析结果与覆盖.xlsx')))
assert wb['交付数据']['D2'].value==0.5 and wb['交付数据']['J2'].value is True and wb['交付数据']['L2'].value==19.5
payload['script']['pages'][0]['key_message']='选择比例为 99.5%'
try:build_package(payload)
except ValueError:pass
else:raise AssertionError('Unbound prose number must be rejected')
print('Weighted base, mean unit and unbound prose number checks passed')
