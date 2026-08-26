"""Regression coverage for multi-level crosstab headers."""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from pptx_report import build_jd_report  # noqa: E402
from pptx_report.build_jd_report import apply_dimension, parse_crosstab  # noqa: E402
from pptx_report.cli import _collect_segments  # noqa: E402


def main() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        path = Path(temp_dir) / "multi-dimension-crosstab.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Table (%)"
        rows = [
            ["CAPTION:[B16].请问您的性别是"],
            [
                None, "Total", "人群", None, None, "使用用户", None, None,
                "目标人群", None, None, None, "购买者年龄", None, None, None,
                "购买者性别", None,
            ],
            [
                None, "Total", "老人组", "儿童组", "老人+儿童组", "老人", "儿童",
                "老人+儿童", "目标人群总体", "老人", "儿童", "老人+儿童",
                "30-34岁", "35-39岁", "40-44岁", "45-49岁", "男性", "女性",
            ],
            ["BASE", 200, 60, 60, 80, 76, 61, 63, 196, 75, 60, 61, 69, 83, 35, 7, 80, 120],
            ["男", 0.4, 0.58, 0.28, 0.35, 0.51, 0.28, 0.38, 0.39, 0.51, 0.27, 0.36, 0.46, 0.22, 0.6, 1, 1, None],
            ["女", 0.6, 0.42, 0.72, 0.65, 0.49, 0.72, 0.62, 0.61, 0.49, 0.73, 0.64, 0.54, 0.78, 0.4, None, None, 1],
        ]
        for row in rows:
            sheet.append(row)
        workbook.save(path)

        questions = parse_crosstab(str(path))
        assert len(questions) == 1
        assert questions[0]["segments"] == ["Total", "老人组", "儿童组", "老人+儿童组"]
        assert list(questions[0]["data"]) == questions[0]["segments"]

        groups = build_jd_report._cached_dimension_groups
        assert [group["name"] for group in groups] == [
            "人群", "使用用户", "目标人群", "购买者年龄", "购买者性别",
        ]
        age_questions = apply_dimension(questions, groups, "购买者年龄")
        assert age_questions[0]["segments"] == [
            "Total", "30-34岁", "35-39岁", "40-44岁", "45-49岁",
        ]
        assert age_questions[0]["data"]["30-34岁"] == [0.46, 0.54]

        pivot_path = Path(temp_dir) / "spss-pivot-crosstab.xlsx"
        pivot_workbook = Workbook()
        pivot_sheet = pivot_workbook.active
        pivot_sheet.title = "Sheet1"
        pivot_rows = [
            [None, None, "省份", None, None, None, "年龄", None, None, None],
            [None, None, "总计", None, "上海", None, "18-24岁", None, "25-34岁", None],
            [None, None, "计数", "列 N %", "计数", "列 N %", "计数", "列 N %", "计数", "列 N %"],
            ["Q1.满意度", "总计", 100, 1, 40, 1, 45, 1, 55, 1],
            [None, "满意", 60, 0.6, 30, 0.75, 31, 0.69, 29, 0.53],
            [None, "不满意", 40, 0.4, 10, 0.25, 14, 0.31, 26, 0.47],
            ["Q2.购买意愿", "总计", 100, 1, 40, 1, 45, 1, 55, 1],
            [None, "愿意", 70, 0.7, 32, 0.8, 34, 0.76, 36, 0.65],
            [None, "不愿意", 30, 0.3, 8, 0.2, 11, 0.24, 19, 0.35],
        ]
        for row in pivot_rows:
            pivot_sheet.append(row)
        pivot_workbook.save(pivot_path)

        pivot_questions = parse_crosstab(str(pivot_path))
        assert len(pivot_questions) == 2
        assert pivot_questions[0]["code"] == "Q1"
        assert pivot_questions[0]["segments"] == ["Total", "上海"]
        pivot_groups = build_jd_report._cached_dimension_groups
        assert [group["name"] for group in pivot_groups] == ["省份", "年龄"]
        pivot_age = apply_dimension(pivot_questions, pivot_groups, "年龄")
        assert pivot_age[0]["segments"] == ["Total", "18-24岁", "25-34岁"]
        assert pivot_age[0]["data"]["18-24岁"] == [0.69, 0.31]
        assert pivot_age[0]["base"]["25-34岁"] == 55
        pivot_summary = _collect_segments(str(pivot_path))
        assert pivot_summary["segments"] == ["Total", "上海", "18-24岁", "25-34岁"]
        assert [group["name"] for group in pivot_summary["dimension_groups"]] == ["省份", "年龄"]

        flat_path = Path(temp_dir) / "flat-question-crosstab.xlsx"
        flat_workbook = Workbook()
        flat_sheet = flat_workbook.active
        flat_sheet.title = "交叉表"
        flat_rows = [
            ["题目/选项", "总体", "C0总体", "C0-高意向", "C1总体", "C1-已购买"],
            ["有效样本量", "n=100", "n=40", "n=20", "n=60", "n=30"],
            ["Q1. Q1.购买意愿【单选】", None, None, None, None, None],
            ["愿意", 0.7, 0.8, 0.9, 0.6, 0.7],
            ["不愿意", 0.3, 0.2, 0.1, 0.4, 0.3],
            ["Q2_1. 使用场景【多选】", None, None, None, None, None],
            ["家庭", 0.6, 0.7, 0.8, 0.5, 0.6],
            ["户外", 0.4, 0.3, 0.2, 0.5, 0.4],
        ]
        for row in flat_rows:
            flat_sheet.append(row)
        flat_workbook.save(flat_path)

        flat_questions = parse_crosstab(str(flat_path))
        assert len(flat_questions) == 2
        assert flat_questions[0]["code"] == "Q1"
        assert flat_questions[0]["title"] == "购买意愿【单选】"
        assert flat_questions[0]["segments"] == [
            "Total", "C0总体", "C0-高意向", "C1总体", "C1-已购买",
        ]
        assert flat_questions[0]["base"]["Total"] == 100
        assert flat_questions[1]["data"]["C1-已购买"] == [0.6, 0.4]
        flat_groups = build_jd_report._cached_dimension_groups
        assert [group["name"] for group in flat_groups] == ["C0", "C1"]
        flat_c0 = apply_dimension(flat_questions, flat_groups, "C0")
        assert flat_c0[0]["segments"] == ["Total", "C0总体", "C0-高意向"]

        exported_path = Path(temp_dir) / "exported-styled-crosstab.xlsx"
        exported_workbook = Workbook()
        directory_sheet = exported_workbook.active
        directory_sheet.title = "目录"
        directory_sheet.append(["交叉表目录"])
        for sheet_name, metric, values in (
            ("频数", "频数", (60, 30)),
            ("百分比", "百分比", (0.6, 0.75)),
        ):
            exported_sheet = exported_workbook.create_sheet(sheet_name)
            exported_sheet.append(["CAPTION:1. Q1. 购买意愿"])
            exported_sheet.append(["", None, "总体", "购买人群"])
            exported_sheet.append(["", None, "总体", "高意向"])
            exported_sheet.append(["", None, metric, metric])
            exported_sheet.append(["BASE", None, 100, 40])
            exported_sheet.append([None, None, None, None])
            exported_sheet.append(["Q1. 购买意愿", None, None, None])
            exported_sheet.append(["愿意", None, values[0], values[1]])
        exported_workbook.save(exported_path)

        exported_questions = parse_crosstab(str(exported_path))
        assert len(exported_questions) == 1
        assert exported_questions[0]["code"] == "1"
        assert exported_questions[0]["title"] == "Q1. 购买意愿"
        assert exported_questions[0]["segments"] == ["总体", "高意向"]
        assert exported_questions[0]["data"]["总体"] == [0.6]
        assert exported_questions[0]["data"]["高意向"] == [0.75]

        adaptive_path = Path(temp_dir) / "adaptive-crosstab.xlsx"
        adaptive_workbook = Workbook()
        count_sheet = adaptive_workbook.active
        count_sheet.title = "Data_A"
        percent_sheet = adaptive_workbook.create_sheet("Data_B")
        for target_sheet, metric, values in (
            (count_sheet, "频数", (60, 30)),
            (percent_sheet, "百分比", ("60%", 75)),
        ):
            target_sheet.append(["CAPTION：[Q9]。购买意愿"])
            target_sheet.append(["", None, None, "全体", "购买人群"])
            target_sheet.append(["", None, None, "全部", "女性"])
            target_sheet.append(["", None, None, metric, metric])
            target_sheet.append(["有效样本量", None, None, 100, 40])
            target_sheet.append(["愿意", None, None, values[0], values[1]])
        adaptive_workbook.save(adaptive_path)

        adaptive_questions = parse_crosstab(str(adaptive_path))
        assert len(adaptive_questions) == 1
        assert adaptive_questions[0]["code"] == "Q9"
        assert adaptive_questions[0]["title"] == "购买意愿"
        assert adaptive_questions[0]["segments"] == ["全部", "女性"]
        assert adaptive_questions[0]["base"] == {"全部": 100, "女性": 40}
        assert adaptive_questions[0]["data"]["全部"] == [0.6]
        assert adaptive_questions[0]["data"]["女性"] == [0.75]
    print("crosstab parser smoke passed")


if __name__ == "__main__":
    main()
