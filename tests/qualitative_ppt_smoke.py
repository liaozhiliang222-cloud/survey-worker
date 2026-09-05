from __future__ import annotations

import base64
import json
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from pptx_report.qualitative_renderer import PAGE_TYPES, prepare_pages, render_qualitative_report
from pptx_report.qualitative_service import (
    OfficeCliUnavailableError,
    QualitativePptRenderService,
)
from pptx_report.officecli_runner import OfficeCliRunner
from deploy.aliyun_api import app


FIXTURE = ROOT / "tests" / "fixtures" / "qualitative-ppt-script.json"
ENTERPRISE_FIXTURE = ROOT / "tests" / "fixtures" / "qualitative-enterprise-template-script.json"
ADAPTIVE_FIXTURE = ROOT / "tests" / "fixtures" / "qualitative-adaptive-stress-script.json"


def _visible_stress_atoms(script: dict) -> list[str]:
    atoms: list[str] = []
    for page in script["pages"]:
        atoms.extend(item["text"] for item in page.get("supporting_findings", []))
        atoms.extend(item["text"] for item in page.get("quotes", []))
        atoms.extend(item["title"] for item in page.get("segments", []))
        atoms.extend(item["label"] for item in page.get("items", []))
        atoms.extend(str(item) for item in page.get("traits", []))
        atoms.extend(str(item) for item in page.get("behaviors", []))
        for block in page.get("content_structure", []):
            atoms.extend(str(item) for item in block.get("items", []))
        atoms.extend(item["segment"] for item in page.get("segment_mapping", []))
    return atoms


def test_realistic_qualitative_deck_is_native_editable() -> None:
    script = json.loads(FIXTURE.read_text(encoding="utf-8"))
    assert len(script["acceptance_project"]["transcript_ids"]) == 12
    assert {page["page_type"] for page in script["pages"]}.issubset(PAGE_TYPES)

    rendered = render_qualitative_report(script)
    assert 10 <= rendered["slide_count"] <= 20
    assert rendered["render_llm_tokens"] == 0
    assert rendered["validation"]["passed"] is True
    assert rendered["object_counts"]["editable_native"] is True
    assert rendered["object_counts"]["pictures"] == 0
    assert rendered["object_counts"]["full_slide_images"] == 0
    assert rendered["object_counts"]["text_shapes"] > 100
    assert rendered["object_counts"]["connectors"] > 0
    with ZipFile(BytesIO(rendered["content"])) as archive:
        arrowheads = sum(
            archive.read(name).count(b"<a:tailEnd")
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
    assert arrowheads >= rendered["object_counts"]["connectors"] - 4

    prs = Presentation(BytesIO(rendered["content"]))
    assert len(prs.slides) == rendered["slide_count"]
    all_text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    for page in script["pages"]:
        for quote in page.get("quotes", []):
            assert quote["text"] in all_text
    assert "evidence_" not in all_text.lower()
    assert "segment_" not in all_text.lower()


def test_density_is_split_without_quote_rewrite() -> None:
    script = json.loads(FIXTURE.read_text(encoding="utf-8"))
    page = json.loads(json.dumps(script["pages"][2]))
    page["id"] = "dense"
    page["supporting_findings"] *= 3
    page["quotes"] *= 4
    page["evidence_ids"] = list(dict.fromkeys(page["evidence_ids"]))
    dense_script = {**script, "pages": [page]}
    rendered = render_qualitative_report(dense_script)
    assert rendered["slide_count"] > 1
    assert any(issue["code"] == "PAGE_TOO_DENSE" and issue["severity"] == "fixed" for issue in rendered["validation"]["issues"])
    all_text = "\n".join(shape.text for slide in Presentation(BytesIO(rendered["content"])).slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
    assert page["quotes"][0]["text"] in all_text


def test_adaptive_layout_preserves_all_visible_source_items() -> None:
    script = json.loads(ADAPTIVE_FIXTURE.read_text(encoding="utf-8"))
    pages, issues = prepare_pages(script)
    assert len(script["pages"]) == 6
    assert len(pages) == 12
    assert all(page["page_number"] == index for index, page in enumerate(pages, 1))
    assert sum(issue["code"] == "PAGE_TOO_DENSE" for issue in issues) == 6
    assert all(issue["severity"] == "fixed" for issue in issues)
    assert max(len(page.get("segments", [])) for page in pages) <= 4
    assert max(len(page.get("items", [])) for page in pages) <= 6
    assert max(len(page.get("traits", [])) for page in pages) <= 4
    assert max(len(page.get("behaviors", [])) for page in pages) <= 4
    assert max(len(page.get("segment_mapping", [])) for page in pages) <= 4
    assert max(len(page.get("content_structure", [])) for page in pages) <= 4

    rendered = render_qualitative_report(script)
    assert rendered["slide_count"] == 12
    assert rendered["layout_adaptations"] == {
        "profile": "qualitative_enterprise_v1_1",
        "input_page_count": 6,
        "output_page_count": 12,
        "continuation_page_count": 6,
        "adapted_source_page_count": 6,
        "fixed_issue_count": 6,
        "verbatim_quotes_preserved": True,
    }
    assert rendered["validation"]["passed"] is True
    assert rendered["object_counts"]["editable_native"] is True
    all_text = "\n".join(
        shape.text
        for slide in Presentation(BytesIO(rendered["content"])).slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    for atom in _visible_stress_atoms(script):
        assert atom in all_text, atom


def test_enterprise_template_uses_technology_blue_without_left_rail() -> None:
    script = json.loads(ENTERPRISE_FIXTURE.read_text(encoding="utf-8"))
    assert script["style_profile"]["forbid_left_vertical_bar"] is True
    enterprise_types = {
        "navigation", "executive_summary", "research_framework", "segmentation_map",
        "persona", "journey", "comparison", "evidence_diagnostic",
        "concept_definition", "needs_pyramid", "priority_matrix", "quote_evidence",
        "problem_reason", "recommendation",
    }
    assert {page["page_type"] for page in script["pages"]} == enterprise_types
    assert enterprise_types.issubset(PAGE_TYPES)

    rendered = render_qualitative_report(script)
    assert rendered["slide_count"] == 15
    assert rendered["validation"]["passed"] is True
    assert not any(
        issue["code"] == "UNSUPPORTED_PAGE_TYPE"
        for issue in rendered["validation"]["issues"]
    )
    assert rendered["object_counts"]["editable_native"] is True
    assert rendered["object_counts"]["pictures"] == 0
    prs = Presentation(BytesIO(rendered["content"]))
    names = {
        shape.name
        for slide in prs.slides
        for shape in slide.shapes
    }
    assert "accent_rule" not in names
    assert "section_accent" not in names
    assert "card_accent" not in names
    assert "semantic_tag" in names
    with ZipFile(BytesIO(rendered["content"])) as archive:
        slide_xml = b"\n".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        )
        notes_xml = b"\n".join(
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")
        )
    assert b"176BFF" in slide_xml
    assert b"0A2A66" in slide_xml
    assert b"Evidence IDs" in notes_xml


def test_officecli_is_primary_enterprise_renderer_when_available() -> None:
    runner = OfficeCliRunner()
    if not runner.is_installed():
        return
    script = json.loads(ENTERPRISE_FIXTURE.read_text(encoding="utf-8"))
    rendered = QualitativePptRenderService(
        officecli=runner, quality_mode="required", engine_mode="auto"
    ).render(script)
    assert rendered["renderer"] == {
        "engine": "officecli",
        "preferred_engine": "officecli",
        "mode": "officecli_native",
        "editable_native": True,
    }
    assert rendered["slide_count"] == 15
    assert rendered["validation"]["passed"] is True
    assert rendered["quality_gate"]["passed"] is True
    assert rendered["quality_gate"]["issues"]["data"]["count"] == 0
    assert rendered["object_counts"]["pictures"] == 0
    assert rendered["object_counts"]["text_shapes"] > 250
    assert rendered["object_counts"]["connectors"] > 0

    prs = Presentation(BytesIO(rendered["content"]))
    names = {shape.name for slide in prs.slides for shape in slide.shapes}
    all_text = "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "accent_rule" not in names
    assert "section_accent" not in names
    assert "card_accent" not in names
    assert "Evidence IDs" not in all_text
    assert "layout_variant_editorial_overview" in names
    assert "layout_variant_journey_curve" in names
    assert "layout_variant_fishbone" in names
    assert "layout_variant_action_roadmap" in names
    for page in script["pages"]:
        for quote in page.get("quotes", []):
            assert quote["text"] in all_text


def test_v2_adapts_legacy_layout_description_blocks_without_placeholders() -> None:
    runner = OfficeCliRunner()
    if not runner.is_installed():
        return
    source = json.loads(ENTERPRISE_FIXTURE.read_text(encoding="utf-8"))
    by_type = {page["page_type"]: json.loads(json.dumps(page)) for page in source["pages"]}
    summary = by_type["executive_summary"]
    summary["supporting_findings"] = ["跨平台决策需要连续信息", "价格承诺必须可以兑现", "售后服务需要前置可见"]
    summary["content_structure"] = [{
        "region": "top",
        "title": "三卡片矩阵：发现 + 置信度标签",
        "items": ["①首个发现｜medium｜Ch2", "②第二发现｜low｜Ch3", "③第三发现｜low｜Ch4"],
    }]
    journey = by_type["journey"]
    journey["content_structure"] = [
        {"title": "三阶段旅程图", "body": "阶段一：内容平台—搜索候选；阶段二：综合电商—比较价格；阶段三：品牌商城—核验权益"},
        {"title": "业务含义", "body": "跨平台衔接"},
    ]
    problem = by_type["problem_reason"]
    problem["content_structure"] = problem["content_structure"][:2]
    problem["supporting_findings"] = ["场景条件：弱光与运动叠加", "结果解释：失败原因不可见", "恢复路径：撤回入口不稳定"]
    recommendation = by_type["recommendation"]
    recommendation["content_structure"] = recommendation["content_structure"][:2]
    recommendation["supporting_findings"] = ["问题绑定：先明确高频摩擦", "动作一：统一反馈语言", "动作二：补齐恢复路径"]
    script = {**source, "pages": [summary, journey, problem, recommendation]}

    rendered = QualitativePptRenderService(
        officecli=runner, quality_mode="required", engine_mode="auto"
    ).render(script)
    all_text = "\n".join(
        shape.text
        for slide in Presentation(BytesIO(rendered["content"])).slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )
    assert "待补充证据" not in all_text
    assert "关键发现 2" not in all_text
    assert "内容平台" in all_text
    assert "综合电商" in all_text
    assert "品牌商城" in all_text
    assert rendered["layout_adaptations"]["profile"] == "qualitative_tech_blue_v2"


def test_qualitative_api_is_project_scoped_and_zero_llm() -> None:
    script = json.loads(FIXTURE.read_text(encoding="utf-8"))
    client = TestClient(app)
    denied = client.post("/api/pptx-report/qualitative-report", headers={"X-Project-Id": "other"}, json={"project_id": "p1", "script": script})
    assert denied.status_code == 403
    response = client.post("/api/pptx-report/qualitative-report", headers={"X-Project-Id": "p1"}, json={"project_id": "p1", "script": script, "require_officecli": True})
    if not OfficeCliRunner().is_installed():
        assert response.status_code == 503
        return
    assert response.status_code == 200
    payload = response.json()
    assert payload["template_id"] == "qualitative_tech_blue_v2"
    assert payload["llm_tokens"] == 0
    assert payload["slide_count"] == 15
    assert payload["validation"]["passed"] is True
    assert payload["object_counts"]["editable_native"] is True
    assert payload["renderer"]["engine"] == "officecli"
    assert payload["quality_gate"]["engine"] == "officecli"
    assert payload["quality_gate"]["status"] in {"passed", "unavailable", "error"}
    assert payload["layout_adaptations"]["input_page_count"] == 15
    assert payload["layout_adaptations"]["output_page_count"] == payload["slide_count"]
    assert len(base64.b64decode(payload["content_base64"])) > 40_000


def test_qualitative_preview_api_uses_real_officecli_png_and_source_scope() -> None:
    runner = OfficeCliRunner()
    if not runner.is_installed():
        return
    script = json.loads(ENTERPRISE_FIXTURE.read_text(encoding="utf-8"))
    source_page = next(page for page in script["pages"] if page["id"] == "tpl02")
    client = TestClient(app)
    denied = client.post(
        "/api/pptx-report/qualitative-preview",
        headers={"X-Project-Id": "other"},
        json={"project_id": "p1", "script": script},
    )
    assert denied.status_code == 403
    response = client.post(
        "/api/pptx-report/qualitative-preview",
        headers={"X-Project-Id": "p1"},
        json={
            "project_id": "p1",
            "script": script,
            "template_id": "qualitative_tech_blue_v2",
            "source_page_id": source_page["id"],
        },
    )
    assert response.status_code == 200
    assert response.headers["X-SurveyKit-Preview-Temporary"] == "true"
    payload = response.json()
    assert payload["schema_version"] == "surveykit.qualitative_preview.v1"
    assert payload["render_scope"] == "source_page"
    assert payload["renderer"] == {
        "engine": "officecli",
        "mode": "temporary_preview",
        "temporary": True,
        "artifact_created": False,
    }
    assert payload["render_llm_tokens"] == 0
    assert payload["rendered_slide_count"] == len(payload["slides"]) >= 1
    assert all(slide["source_page_id"] == source_page["id"] for slide in payload["slides"])
    first = payload["slides"][0]
    assert base64.b64decode(first["thumbnail_base64"]).startswith(b"\x89PNG\r\n\x1a\n")
    assert first["thumbnail_mime_type"] == "image/png"
    assert first["page_type"] == "executive_summary"
    assert first["layout_variant"] == "editorial_overview"
    assert {item["id"] for item in first["layout_candidates"]} == {
        "north_star_stack",
        "editorial_overview",
    }
    assert first["evidence"]["count"] == len(source_page["evidence_ids"])
    assert "status" in first["density"]
    assert isinstance(first["validation_issues"], list)

    source_page["layout_variant"] = "north_star_stack"
    source_page.setdefault("layout_spec", {})["variant"] = "north_star_stack"
    changed = client.post(
        "/api/pptx-report/qualitative-preview",
        headers={"X-Project-Id": "p1"},
        json={
            "project_id": "p1",
            "script": script,
            "template_id": "qualitative_tech_blue_v2",
            "source_page_id": source_page["id"],
        },
    )
    assert changed.status_code == 200
    changed_payload = changed.json()
    assert changed_payload["script_fingerprint"] != payload["script_fingerprint"]
    assert changed_payload["rendered_slide_count"] == len(changed_payload["slides"])
    assert all(item["source_page_id"] == "tpl02" for item in changed_payload["slides"])
    assert changed_payload["slides"][0]["layout_variant"] == "north_star_stack"
    assert next(item for item in changed_payload["slides"][0]["layout_candidates"] if item["selected"])["id"] == "north_star_stack"


def test_qualitative_template_catalog_and_preflight_contract() -> None:
    client = TestClient(app)
    response = client.get("/api/pptx-report/qualitative-templates")
    assert response.status_code == 200
    payload = response.json()
    assert payload["default_template_id"] == "qualitative_tech_blue_v2"
    assert payload["renderer"]["preferred_generation_engine"] == "officecli"
    assert isinstance(payload["renderer"]["officecli_installed"], bool)
    template = payload["templates"][0]
    assert template["template_id"] == payload["default_template_id"]
    assert template["name"] == "Tech Blue V2"
    assert template["editable"] is True
    assert template["aspect_ratio"] == "16:9"
    assert template["constraints"]["forbid_left_vertical_bar"] is True
    assert template["constraints"]["native_objects_only"] is True
    assert len(template["layouts"]) == 16
    assert template["preview_layout_ids"] == [
        "editorial_overview",
        "profile_evidence",
        "journey_curve",
        "impact_frequency",
        "fishbone",
        "action_roadmap",
    ]

    script = json.loads(FIXTURE.read_text(encoding="utf-8"))
    rejected = client.post(
        "/api/pptx-report/qualitative-report",
        headers={"X-Project-Id": "p1"},
        json={
            "project_id": "p1",
            "script": script,
            "template_id": "unknown_template",
        },
    )
    assert rejected.status_code == 400
    assert "不支持的定性 PPT 模板" in rejected.json()["error"]["message"]

    required = client.post(
        "/api/pptx-report/qualitative-report",
        headers={"X-Project-Id": "p1"},
        json={
            "project_id": "p1",
            "script": script,
            "template_id": payload["default_template_id"],
            "require_officecli": True,
        },
    )
    if OfficeCliRunner().is_installed():
        assert required.status_code == 200
        assert required.json()["renderer"]["engine"] == "officecli"
    else:
        assert required.status_code == 503
        assert required.json()["error"]["code"] == "OFFICECLI_UNAVAILABLE"


class MissingOfficeCli:
    binary = "/missing/officecli"

    def is_installed(self) -> bool:
        return False

    def probe(self) -> dict:
        return {"installed": False, "version": "", "binary": self.binary}


def test_officecli_quality_gate_policy_is_explicit() -> None:
    script = json.loads(FIXTURE.read_text(encoding="utf-8"))
    preferred = QualitativePptRenderService(
        officecli=MissingOfficeCli(), quality_mode="prefer"
    ).render(script)
    assert preferred["renderer"]["engine"] == "python-pptx"
    assert preferred["renderer"]["fallback"] == {
        "from": "officecli",
        "reason": "unavailable",
    }
    assert preferred["quality_gate"] == {
        "engine": "officecli",
        "mode": "prefer",
        "status": "unavailable",
        "passed": None,
    }
    try:
        QualitativePptRenderService(
            officecli=MissingOfficeCli(), quality_mode="required"
        ).render(script)
    except OfficeCliUnavailableError:
        pass
    else:
        raise AssertionError("required OfficeCLI mode must not silently fall back")
    try:
        QualitativePptRenderService(
            officecli=MissingOfficeCli(), quality_mode="prefer", engine_mode="officecli"
        ).render(script)
    except OfficeCliUnavailableError:
        pass
    else:
        raise AssertionError("explicit OfficeCLI engine mode must not silently fall back")


if __name__ == "__main__":
    test_realistic_qualitative_deck_is_native_editable()
    test_density_is_split_without_quote_rewrite()
    test_adaptive_layout_preserves_all_visible_source_items()
    test_enterprise_template_uses_technology_blue_without_left_rail()
    test_officecli_is_primary_enterprise_renderer_when_available()
    test_v2_adapts_legacy_layout_description_blocks_without_placeholders()
    test_qualitative_api_is_project_scoped_and_zero_llm()
    test_qualitative_preview_api_uses_real_officecli_png_and_source_scope()
    test_qualitative_template_catalog_and_preflight_contract()
    test_officecli_quality_gate_policy_is_explicit()
    output = ROOT / "test-results" / "qualitative-ppt-acceptance.pptx"
    output.parent.mkdir(parents=True, exist_ok=True)
    script = json.loads(FIXTURE.read_text(encoding="utf-8"))
    output.write_bytes(render_qualitative_report(script)["content"])
    enterprise_output = ROOT / "test-results" / "SurveyKit-qualitative-enterprise-template-python-fallback.pptx"
    enterprise_script = json.loads(ENTERPRISE_FIXTURE.read_text(encoding="utf-8"))
    enterprise_output.write_bytes(render_qualitative_report(enterprise_script)["content"])
    runner = OfficeCliRunner()
    acceptance_service = QualitativePptRenderService(
        officecli=runner,
        quality_mode="required" if runner.is_installed() else "prefer",
        engine_mode="auto",
    )
    realistic_rendered = acceptance_service.render(script)
    realistic_output = ROOT / "test-results" / "SurveyKit-qualitative-realistic-officecli-v1.pptx"
    realistic_output.write_bytes(realistic_rendered["content"])
    adaptive_script = json.loads(ADAPTIVE_FIXTURE.read_text(encoding="utf-8"))
    adaptive_rendered = acceptance_service.render(adaptive_script)
    adaptive_output = ROOT / "test-results" / "SurveyKit-qualitative-adaptive-officecli-v1.pptx"
    adaptive_output.write_bytes(adaptive_rendered["content"])
    (ROOT / "test-results" / "SurveyKit-qualitative-adaptive-officecli-v1.json").write_text(
        json.dumps(
            {
                "renderer": adaptive_rendered["renderer"],
                "slide_count": adaptive_rendered["slide_count"],
                "validation": adaptive_rendered["validation"],
                "quality_gate": adaptive_rendered["quality_gate"],
                "object_counts": adaptive_rendered["object_counts"],
                "layout_adaptations": adaptive_rendered["layout_adaptations"],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(output)
    print(enterprise_output)
    print(realistic_output)
    print(adaptive_output)
