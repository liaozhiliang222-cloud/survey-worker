"""Template contract, content fidelity, API and real OfficeCLI regression checks."""
from copy import deepcopy
from io import BytesIO
import base64
import hashlib
import json
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from fastapi.testclient import TestClient
from pptx import Presentation
from deploy.aliyun_api import app
from pptx_report.business_blue_layouts import BUSINESS_BLUE_ID, business_blue_catalog, resolve_business_blue
from pptx_report.business_blue_content import cards_for_page
from pptx_report.qualitative_layouts import layout_candidates
from pptx_report.qualitative_renderer import prepare_pages, render_qualitative_report
from pptx_report.officecli_qualitative_renderer import build_officecli_commands, render_officecli_qualitative_report
from pptx_report.officecli_runner import OfficeCliRunner


def fixture():
    return json.loads((ROOT / "tests/fixtures/business-blue-script.json").read_text(encoding="utf-8"))


def visible_commands(script):
    pages, issues = prepare_pages(script)
    commands = build_officecli_commands(script, pages)
    return pages, issues, "\n".join(c.get("props", {}).get("text", "") for c in commands if c.get("type") == "shape")


def test_registry_and_source():
    catalog = business_blue_catalog()
    assert len(catalog["layouts"]) == len({x["id"] for x in catalog["layouts"]}) == 24
    source = ROOT / "pptx_report/templates/research-business-blue-v1/reference.pptx"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == catalog["source_sha256"]
    for layout in catalog["layouts"]:
        assert any(x["id"] == layout["id"] for x in layout_candidates(layout["page_types"][0], layout["id"], BUSINESS_BLUE_ID))
    assert resolve_business_blue({"page_type": "segmentation_map", "segments": [{"title": "未评分人群"}]}) == "bb_persona_pair"


def test_density_quotes_and_toc_are_lossless():
    script = fixture()
    page = deepcopy(next(p for p in script["pages"] if p["layout_variant"] == "bb_findings"))
    page["content_structure"] = [{"title": f"主题 {i}", "body": (f"保留{i}号内容，" * 35)} for i in range(11)]
    raw_quote = "这是必须逐字保留的原话，'low' 和 interpretation 不能被翻译。" * 40
    page["quotes"] = [{"text": raw_quote, "evidence_id": "demo_evidence", "segment_id": "demo_segment"}]
    script["pages"] = [page]
    original = deepcopy(script)
    prepared, issues, visible = visible_commands(script)
    assert script == original
    assert len(prepared) > 1 and issues
    original_bodies = "".join(c["body"] for c in cards_for_page(page))
    assert "".join(c["body"] for p in prepared for c in p["_bb_cards"]) == original_bodies
    assert "".join(q["text"] for p in prepared for q in p["quotes"]) == raw_quote
    assert "'low'" in visible and "interpretation" in visible
    assert all(p["evidence_ids"] == page["evidence_ids"] for p in prepared)
    toc_pages, _, _ = visible_commands(fixture())
    toc = [c for p in toc_pages if p["layout_variant"] == "bb_contents" for c in p["_bb_toc"]]
    for c in toc:
        first = next(p for p in toc_pages if p.get("chapter") == c["title"] and p["layout_variant"] not in {"bb_cover", "bb_contents"})
        assert c["body"] == f"第 {first['page_number']} 页"
    assert all(len(p.get("_bb_toc", [])) <= 6 for p in toc_pages)


def test_invalid_templates_and_data_fail_explicitly():
    script = fixture()
    for mutation in (lambda p: p.update(layout_variant="bb_missing"), lambda p: p.update(data_points=[{"value": .5}]),
                     lambda p: p.update(layout_variant="bb_quotes", page_type="quote_evidence", quotes=[])):
        s = deepcopy(script); s["pages"] = [deepcopy(s["pages"][0])]; mutation(s["pages"][0])
        try: prepare_pages(s)
        except ValueError: pass
        else: raise AssertionError("invalid input silently accepted")
    try: render_qualitative_report(script)
    except ValueError as exc: assert "OfficeCLI" in str(exc)
    else: raise AssertionError("template silently rendered with legacy style")


def test_native_and_api():
    runner = OfficeCliRunner()
    assert runner.is_installed(), "OfficeCLI required for the phase-one acceptance test"
    script = fixture()
    rendered = render_officecli_qualitative_report(script, runner)
    assert rendered["generation_quality_gate"]["passed"] and rendered["validation"]["passed"]
    assert rendered["object_counts"]["full_slide_images"] == 0
    prs = Presentation(BytesIO(rendered["content"]))
    assert len(prs.slides) == rendered["slide_count"]
    visible = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    assert "400-000-0000" not in visible and "15.8" not in visible
    for p in rendered["prepared_pages"]:
        for c in p["_bb_cards"]:
            if p["layout_variant"] != "bb_contents": assert c["body"] in visible
        for q in p["quotes"]: assert q["text"] in visible
    output = ROOT / "tests/output/business-blue"
    output.mkdir(parents=True, exist_ok=True)
    (output / "all-layouts.pptx").write_bytes(rendered["content"])
    client = TestClient(app)
    catalogs = client.get("/api/pptx-report/qualitative-templates").json()
    assert {c["template_id"] for c in catalogs["templates"]} == {BUSINESS_BLUE_ID, "qualitative_tech_blue_v2"}
    tiny = deepcopy(script); tiny["pages"] = [script["pages"][0]]
    payload = {"project_id": "blue_test", "template_id": BUSINESS_BLUE_ID, "script": tiny}
    headers = {"X-Project-Id": "blue_test"}
    preview = client.post("/api/pptx-report/qualitative-preview", headers=headers, json=payload)
    assert preview.status_code == 200, preview.text
    p = preview.json()
    assert base64.b64decode(p["slides"][0]["thumbnail_base64"]).startswith(b"\x89PNG")
    assert p["slides"][0]["layout_variant"] == "bb_cover"
    response = client.post("/api/pptx-report/qualitative-report", headers=headers, json=payload)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["template_version"] == "1.0.0" and result["template_name"] == "蓝色商务研究报告"
    assert result["quality_gate"]["passed"] and result["slide_count"] == p["slide_count"]
    wrong_scope = client.post("/api/pptx-report/qualitative-report", headers={"X-Project-Id": "other"}, json=payload)
    assert wrong_scope.status_code == 403
    stale = deepcopy(payload); stale["script"]["style_profile"]["version"] = "0.0.0"
    assert client.post("/api/pptx-report/qualitative-report", headers=headers, json=stale).status_code == 400
    os.environ["SURVEYKIT_BUSINESS_BLUE_ENABLED"] = "0"
    try:
        assert len(client.get("/api/pptx-report/qualitative-templates").json()["templates"]) == 1
        assert client.post("/api/pptx-report/qualitative-report", headers=headers, json=payload).status_code == 400
    finally: os.environ.pop("SURVEYKIT_BUSINESS_BLUE_ENABLED", None)


if __name__ == "__main__":
    for test in (test_registry_and_source, test_density_quotes_and_toc_are_lossless, test_invalid_templates_and_data_fail_explicitly, test_native_and_api):
        test(); print(test.__name__ + ": passed", flush=True)
