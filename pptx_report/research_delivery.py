"""Bounded mixed research delivery: native charts/text plus auditable workbook."""
from copy import deepcopy
from io import BytesIO
import json
import math
import re
import zipfile
from collections import defaultdict
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from openpyxl import Workbook
from .qualitative_renderer import _title, _footer, _textbox, _card, _theme, _add_speaker_notes, W, H, NAVY, BLUE, MUTED
from .resource_gate import serialized_heavy_task

TYPES = {'cover','section_intro','data_insight','comparison','quote_evidence','qualitative_insight','qualitative_summary','summary','executive_summary','recommendation'}

def number(value):
    if isinstance(value, bool) or value is None or value == '':
        raise ValueError('缺失数值不能作为 0 交付')
    value = float(str(value).removesuffix('%'))
    if not math.isfinite(value):
        raise ValueError('无效数值')
    return value

def check(payload):
    if payload.get('schema_version') != 'surveykit.report_delivery.v1' or not payload.get('can_export'):
        raise ValueError('请先完成来源和结论复核')
    pages = payload.get('script', {}).get('pages', [])
    if not pages or len(pages) > 80:
        raise ValueError('报告需包含 1–80 页')
    records = {r['id']:r for r in payload.get('records', [])}
    for page in pages:
        if page.get('page_type') not in TYPES:
            raise ValueError('当前交付不支持此页型')
        if len(page.get('title', '')) > 160:
            raise ValueError('标题超过 160 字，请先精简标题；正文可自动续页')
        ids = set(page.get('evidence_ids', []) + [x.get('evidence_id') for x in page.get('data_points', []) + page.get('quotes', [])])
        if page.get('page_type') not in {'cover','section_intro'} and not ids:
            raise ValueError('缺证据结论仍为待验证')
        if page.get('delivery_review_pending'):
            raise ValueError('更新后的结论仍待复核')
        for eid in ids:
            r=records.get(eid)
            if not r or r.get('unavailable') or page.get('delivery_snapshot', {}).get(eid) != r.get('fingerprint'):
                raise ValueError('证据缺失或来源快照不一致')
        allowed=set()
        def collect(v):
            if isinstance(v,dict):
                for x in v.values():collect(x)
            elif isinstance(v,list):
                for x in v:collect(x)
            else:
                try:allowed.add(number(v))
                except (ValueError,TypeError):pass
        for eid in ids:
            value=records[eid]['evidence']['value'];collect(value)
            allowed.update(float(n) for n in re.findall(r'-?\d+(?:\.\d+)?',value.get('quote','')))
        prose=[page.get('title',''),page.get('key_message',''),page.get('subtitle',''),*page.get('supporting_points',[])]
        for b in page.get('content_structure',[]):prose.extend([b.get('title',''),b.get('body',''),*b.get('items',[])])
        text=re.sub(r'\b[QPV]\d+\b','',' '.join(str(x) for x in prose))
        if any(float(n) not in allowed for n in re.findall(r'-?\d+(?:\.\d+)?',text)):
            raise ValueError('结论包含未匹配源证据的数字')
        for point in page.get('data_points', []):
            r=records[point['evidence_id']];e=r['evidence'];field=point.get('value_field') or 'value'
            if e['value'].get('unit') and point.get('unit') != e['value']['unit']:raise ValueError('数据单位与源证据不一致')
            if field not in {'value','total','overall','base','total_base','mean','nps','percent'} or number(point['value']) != number(e['value'].get(field)) or point.get('data_source_id',e['source_id']) != e['source_id']:
                raise ValueError('数值与来源不一致')
            if len(point.get('label','')) > 100:
                raise ValueError('图表标签超过 100 字，请精简标签')
            if page.get('visual_spec',{}).get('type') not in {None,'none','text_summary','bar_chart','comparison','table'}:
                raise ValueError('交付首批定量页支持横向对比图，请先调整图表类型')
        for q in page.get('quotes', []):
            if not q.get('text') or q['text'] not in records[q['evidence_id']]['evidence']['value'].get('quote',''):
                raise ValueError('原声无法逐字回溯')
    return records

def chunks(text, size):
    parts=[];start=0;lines=0
    for i,char in enumerate(text):
        if char=='\n':lines+=1
        if i-start+1>=size or lines>=5:
            parts.append(text[start:i+1]);start=i+1;lines=0
    if start<len(text):parts.append(text[start:])
    return parts

def prepare(script):
    result=[]
    for page in script['pages']:
        parts=[]
        # No incompatible units share an axis; each group has explicit source units.
        groups=defaultdict(list)
        for point in page.get('data_points',[]):groups[point.get('unit','')].append(point)
        for unit,points in groups.items():
            for i in range(0,len(points),6):parts.append(('chart',points[i:i+6]))
        for q in page.get('quotes',[]):
            for part in chunks(q['text'],220):parts.append(('quote',{**q,'text':part}))
        paragraphs=[]
        for text in [page.get('key_message','') if not parts or len(page.get('key_message',''))>130 else '',page.get('subtitle',''),*page.get('supporting_points',[])]:
            if text:paragraphs.extend(chunks(str(text),180))
        for block in page.get('content_structure',[]):
            text='\n'.join(str(t) for t in [block.get('title',''),block.get('body',''),*block.get('items',[])] if t)
            paragraphs.extend(chunks(text,180))
        for i in range(0,len(paragraphs),2):parts.append(('text',paragraphs[i:i+2]))
        if not parts:parts=[('text',[page.get('purpose','') or '研究报告'])]
        for index,(kind,content) in enumerate(parts):
            result.append((deepcopy(page),kind,content,index+1,len(parts)))
    if len(result)>400:raise ValueError('自动续页超过 400 页，请拆分报告')
    return result

def cell(v):
    if v is None:return '未提供'
    if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False)
    # Prevent spreadsheet formulas from user-provided labels.
    if isinstance(v,str) and v[:1] in '=+-@':v="'"+v
    return v

def workbook(payload, records):
    wb=Workbook();ws=wb.active;ws.title='交付数据'
    ws.append(['脚本页ID','标题','标签','原始值','显示值','单位','Evidence ID','结果ID','数据版本','加权状态','原始base','加权base','缺失规则'])
    for page in payload['script']['pages']:
        for p in page.get('data_points',[]):
            r=records[p['evidence_id']];v=r['evidence']['value'];dataset=r.get('dataset') or {};weighted=dataset.get('type')=='weighted' if dataset else v.get('weighted')
            ws.append([cell(x) for x in [page['id'],page['title'],p['label'],number(p['value']),f"{number(p['value']):.1f}{p.get('unit','')}",p.get('unit',''),r['id'],r['evidence']['source_id'],dataset.get('id'),weighted,v.get('raw_base',v.get('base')),v.get('weighted_base'),v.get('missing_policy') or (dataset.get('metadata') or {}).get('missing_policy')]])
    quotes=wb.create_sheet('原声');quotes.append(['页ID','Evidence ID','访谈ID','片段ID','逐字原声'])
    for page in payload['script']['pages']:
        for q in page.get('quotes',[]):
            v=records[q['evidence_id']]['evidence']['value'];quotes.append([cell(x) for x in [page['id'],q['evidence_id'],v.get('transcript_id'),v.get('segment_id'),q['text']]])
    cover=wb.create_sheet('题目覆盖');cover.append(['结果ID','字段','分析状态','分析原因','报告引用状态','报告页ID'])
    for row in payload.get('coverage',[]):cover.append([cell(row.get(k)) for k in ['result_id','variable','status','reason']]+['已引用' if row.get('report_page_ids') else '未引用',','.join(row.get('report_page_ids',[]))])
    if not payload.get('coverage'):cover.append(['未提供全题清单','','仅声明所引用证据覆盖，不声明全问卷覆盖',''])
    source_sheet=wb.create_sheet('源分析结果');source_sheet.append(['结果ID','题目','分群字段','选项','人群','值','单位','原始人数','原始base','加权人数','加权base'])
    seen=set()
    for r in records.values():
        source=r.get('source') or {}
        if source.get('id') in seen:continue
        seen.add(source.get('id'))
        for item in (source.get('result') or {}).get('results',[]):
            if item.get('metric')=='distribution':
                for category in item.get('categories',[]):
                    for g in category.get('groups',[]):source_sheet.append([cell(x) for x in [source.get('id'),item.get('variable'),item.get('banner'),category.get('category'),g.get('segment'),g.get('percent'),'%',g.get('raw_count',g.get('count')),g.get('raw_base',g.get('base')),g.get('weighted_count'),g.get('weighted_base')]])
            else:
                for g in item.get('groups',[]):source_sheet.append([cell(x) for x in [source.get('id'),item.get('variable'),item.get('banner'),'',g.get('segment'),g.get('value'),item.get('metric'),None,g.get('base'),None,g.get('weighted_base')]])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=24
    out=BytesIO();wb.save(out);return out.getvalue()

@serialized_heavy_task
def build_package(payload):
    records=check(payload);script=payload['script'];prepared=prepare(script)
    prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H);theme=_theme({})
    for owner in [*prs.slide_masters,*prs.slide_layouts]:
        for shape in list(owner.shapes):
            if shape.is_placeholder and shape.placeholder_format.type in {13,15,16}:shape._element.getparent().remove(shape._element)
    manifest_pages=[]
    for number_,(page,kind,content,part,total) in enumerate(prepared,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);page['page_number']=number_
        # The title region is intentionally taller than legacy layouts for long titles.
        title=page['title'];_textbox(slide,.62,.45,12.05,1.30,title,size=22 if len(title)<65 else 17,color=NAVY,bold=True,theme=theme,margin=0,name='slide_title')
        _textbox(slide,.62,1.78,12,.40,(page.get('key_message','') if len(page.get('key_message',''))<=130 else '') + (f'（续 {part}/{total}）' if part>1 else ''),size=10,color=MUTED,theme=theme)
        if kind=='chart':
            data=CategoryChartData();data.categories=[p['label'] for p in content];data.add_series(content[0].get('unit') or '数值',[number(p['value']) for p in content])
            chart=slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(.65),Inches(2.18),Inches(11.95),Inches(3.85),data).chart
            chart.has_legend=False;chart.has_title=False;chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=RGBColor.from_string(BLUE);plot=chart.plots[0];plot.has_data_labels=True;plot.data_labels.position=XL_DATA_LABEL_POSITION.OUTSIDE_END;plot.data_labels.font.size=Pt(12)
            unit=content[0].get('unit','');_textbox(slide,.72,2.05,3,.22,'单位：'+(unit or '源证据数值'),size=10,color=MUTED,theme=theme);plot.data_labels.number_format='0.0"%"' if unit=='%' else '0.0';plot.data_labels.number_format_is_linked=False
            chart.category_axis.tick_labels.font.size=Pt(11);chart.value_axis.tick_labels.font.size=Pt(10)
            if unit=='%':chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=110
            bases=[]
            for p in content:
                r=records[p['evidence_id']];v=r['evidence']['value'];ds=r.get('dataset') or {};base=v.get('raw_base',v.get('base'));weighted=ds.get('type')=='weighted' if ds else v.get('weighted')
                bases.append(f"{p['label']}：base={base if base is not None else '未提供'}"+('（小样本）' if isinstance(base,(int,float)) and base<30 else '')+('；已加权' if weighted else '；未加权' if weighted is False else '；加权状态未记录'))
            _textbox(slide,.7,6.10,11.9,.70,'；'.join(bases),size=9,color=MUTED,theme=theme)
        elif kind=='quote':
            _card(slide,.72,2.30,11.85,3.85,content.get('source_label') or '访谈原声',[content['text']],theme=theme,body_size=19)
        else:
            for i,text in enumerate(content):_card(slide,.72,2.22+i*2.23,11.85,2.08,'结论与依据' if i==0 else '补充信息',[text],theme=theme,body_size=13)
        page['source_notes']='证据、来源版本及统计口径见随包清单';_footer(slide,page,theme);_add_speaker_notes(slide,page)
        manifest_pages.append({'slide':number_,'source_page_id':page['id'],'part':part,'parts':total,'kind':kind})
    out=BytesIO();prs.save(out);ppt=out.getvalue()
    manifest={'schema_version':'surveykit.report_package.v1','project_id':payload['project_id'],'artifact_id':payload['artifact']['id'],'artifact_version':payload['artifact'].get('version'),'captured_at':payload['captured_at'],'pages':manifest_pages,'sources':[{'evidence_id':r['id'],'source_id':r['evidence']['source_id'],'source_type':r['evidence']['source_type'],'fingerprint':r['fingerprint'],'dataset':r.get('dataset'),'source_version':(r.get('source') or {}).get('version'),'transcript_version':(r.get('source') or {}).get('transcript_version'),'segment_id':r['evidence']['value'].get('segment_id')} for r in records.values()]}
    notes='# 数据口径与交付说明\n\n数值逐点取自所绑定证据，显示保留 1 位小数；Excel 同时保留原始精度。百分比按证据既有单位输出，不自动乘 100。\n\n缺失值不补零；未知 base、权重或缺失规则明确标记未提供。小样本（base<30）仅供方向性参考。加权结果为描述统计，不宣称显著性或因果关系。\n\n题目覆盖见 Excel，未提供全题清单时不声明完整问卷覆盖。自动续页不改变引用原文。\n\n本包冻结于来源清单的时间与版本，后续源数据变化需重新检查并生成。\n'
    archive=BytesIO()
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('研究报告.pptx',ppt);z.writestr('分析结果与覆盖.xlsx',workbook(payload,records));z.writestr('数据口径说明.md',notes);z.writestr('来源版本清单.json',json.dumps(manifest,ensure_ascii=False,indent=2));z.writestr('题目覆盖清单.json',json.dumps(payload.get('coverage',[]),ensure_ascii=False,indent=2))
    return {'content':archive.getvalue(),'pptx':ppt,'manifest':manifest,'slide_count':len(prepared)}
